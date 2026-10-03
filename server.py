#!/usr/bin/env python3
"""android-lan-file-server: local file browser + in-browser preview for Termux Ubuntu."""

from __future__ import annotations

import argparse
import concurrent.futures
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile
import threading
import time
import zipfile
from pathlib import Path
from urllib.parse import quote

from flask import (
    Flask,
    abort,
    jsonify,
    render_template,
    request,
    send_file,
)

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

APP_ROOT = Path(__file__).resolve().parent
STATIC_DIR = APP_ROOT / "static"
TEMPLATE_DIR = APP_ROOT / "templates"

_TRUE_VALUES = {"1", "true", "yes", "on", "y"}


def _parse_conf_file(path: Path) -> dict[str, str]:
    data: dict[str, str] = {}
    if not path.is_file():
        return data
    try:
        text = path.read_text(encoding="utf-8", errors="ignore")
    except OSError:
        return data
    for line in text.splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if "=" not in line:
            continue
        key, val = line.split("=", 1)
        key = key.strip()
        val = val.strip().strip('"').strip("'")
        if key:
            data[key] = val
    return data


def _resolve_conf_path() -> Path:
    env = os.environ.get("ANDROID_LAN_CONF")
    if env:
        return Path(env).expanduser()
    for cand in (
        APP_ROOT / "android-lan-file-server.conf",
        Path("/opt/android-lan-file-server/android-lan-file-server.conf"),
        Path("/etc/android-lan-file-server/android-lan-file-server.conf"),
    ):
        if cand.is_file():
            return cand
    return APP_ROOT / "android-lan-file-server.conf"


CONF_PATH = _resolve_conf_path()
CONF = _parse_conf_file(CONF_PATH)


def _resolve_local_conf_path() -> Path:
    env = os.environ.get("ANDROID_LAN_LOCAL_CONF")
    if env:
        return Path(env).expanduser()
    return APP_ROOT / "android-lan-file-server.local.conf"


# 本地私有配置（gitignore 不入库），优先级高于 android-lan-file-server.conf
LOCAL_CONF = _parse_conf_file(_resolve_local_conf_path())


def _cfg(env_key: str, conf_key: str, default: str) -> str:
    """Env var > local private conf > android-lan-file-server.conf > default."""
    if env_key in os.environ and os.environ[env_key] != "":
        return os.environ[env_key]
    local = LOCAL_CONF.get(conf_key, "")
    if local != "":
        return local
    return CONF.get(conf_key, default)


def _cfg_bool(env_key: str, conf_key: str, default: str = "no") -> bool:
    return _cfg(env_key, conf_key, default).strip().lower() in _TRUE_VALUES


DEFAULT_ROOT = _cfg("ANDROID_LAN_ROOT", "DEFAULT_ROOT", "/sdcard")
ALLOW_OUTSIDE_ROOT = _cfg_bool(
    "ANDROID_LAN_ALLOW_OUTSIDE_ROOT", "ALLOW_OUTSIDE_ROOT", "no"
)
CACHE_DIR = Path(
    _cfg("ANDROID_LAN_CACHE", "ANDROID_LAN_CACHE", "/tmp/android-lan-file-server-cache")
)
ANDROID_LAN_ROOT = Path(DEFAULT_ROOT).expanduser().resolve()
DEFAULT_HOST = _cfg("ANDROID_LAN_HOST", "ANDROID_LAN_HOST", "127.0.0.1")
DEFAULT_PORT = int(_cfg("ANDROID_LAN_PORT", "ANDROID_LAN_PORT", "9523"))
MAX_TEXT_BYTES = int(os.environ.get("ANDROID_LAN_MAX_TEXT", str(2 * 1024 * 1024)))
MAX_LIST_ENTRIES = int(os.environ.get("ANDROID_LAN_MAX_LIST", "5000"))
# 版本号单一来源：/healthz 与首页「关于」弹窗共用
APP_VERSION = "1.1.0"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")

app = Flask(
    __name__,
    static_folder=str(STATIC_DIR),
    template_folder=str(TEMPLATE_DIR),
)
app.config["MAX_CONTENT_LENGTH"] = 64 * 1024 * 1024
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0

_pptx_lock = threading.Lock()

TEXT_EXTS = {
    ".txt", ".log", ".ini", ".conf", ".cfg", ".json", ".xml", ".yml", ".yaml",
    ".toml", ".csv", ".tsv", ".sh", ".bash", ".zsh", ".py", ".js", ".ts",
    ".html", ".htm", ".css", ".c", ".cpp", ".h", ".java", ".go", ".rs",
    ".sql", ".bat", ".ps1", ".env", ".md5", ".gitignore",
}
MD_EXTS = {".md", ".markdown", ".mdown", ".mkd", ".mkdn"}
IMAGE_EXTS = {
    ".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".svg", ".ico", ".avif",
}
AUDIO_EXTS = {
    ".mp3", ".m4a", ".aac", ".flac", ".wav", ".ogg", ".oga", ".opus", ".amr",
}
VIDEO_EXTS = {
    ".mp4", ".webm", ".mkv", ".mov", ".m4v", ".avi", ".3gp", ".ts", ".flv",
}
PDF_EXTS = {".pdf"}
PPT_EXTS = {".ppt", ".pptx", ".pps", ".ppsx", ".odp"}
OFFICE_EXTS = {".doc", ".docx", ".odt", ".rtf", ".xls", ".xlsx", ".ods"}
EPUB_EXTS = {".epub"}

PREVIEW_KINDS = {
    "pdf", "ppt", "office", "markdown", "text", "audio", "video", "image",
    "epub",
}

MIME_OVERRIDES = {
    ".md": "text/markdown; charset=utf-8",
    ".markdown": "text/markdown; charset=utf-8",
    ".js": "application/javascript; charset=utf-8",
    ".mjs": "application/javascript; charset=utf-8",
    ".ts": "text/plain; charset=utf-8",
    ".json": "application/json; charset=utf-8",
    ".yml": "text/plain; charset=utf-8",
    ".yaml": "text/plain; charset=utf-8",
    ".csv": "text/plain; charset=utf-8",
    ".svg": "image/svg+xml",
    ".webm": "video/webm",
    ".m4v": "video/mp4",
    ".m4a": "audio/mp4",
    ".opus": "audio/ogg",
    ".flac": "audio/flac",
    ".ogv": "video/ogg",
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _fmt_size(n: int) -> str:
    if n < 0:
        return "-"
    units = ["B", "KB", "MB", "GB", "TB"]
    size = float(n)
    for u in units:
        if size < 1024 or u == units[-1]:
            if u == "B":
                return f"{int(size)} {u}"
            return f"{size:.1f} {u}"
        size /= 1024
    return f"{n} B"


def _fmt_time(ts: float) -> str:
    try:
        return time.strftime("%Y-%m-%d %H:%M", time.localtime(ts))
    except (OverflowError, OSError, ValueError):
        return "-"


def _ext(path: Path) -> str:
    return path.suffix.lower()


def detect_kind(path: Path) -> str:
    ext = _ext(path)
    if ext in PDF_EXTS:
        return "pdf"
    if ext in PPT_EXTS:
        return "ppt"
    if ext in OFFICE_EXTS:
        return "office"
    if ext in MD_EXTS:
        return "markdown"
    if ext in AUDIO_EXTS:
        return "audio"
    if ext in VIDEO_EXTS:
        return "video"
    if ext in IMAGE_EXTS:
        return "image"
    if ext in EPUB_EXTS:
        return "epub"
    if ext in TEXT_EXTS:
        return "text"
    try:
        if path.is_file() and path.stat().st_size <= MAX_TEXT_BYTES:
            with path.open("rb") as f:
                head = f.read(4096)
            if b"\x00" in head:
                return "binary"
            try:
                head.decode("utf-8")
                return "text"
            except UnicodeDecodeError:
                return "binary"
    except OSError:
        pass
    return "binary"


def _rel_for_url(path: Path) -> str:
    """Build URL path fragment for a filesystem path."""
    try:
        return path.relative_to(ANDROID_LAN_ROOT).as_posix()
    except ValueError:
        return "@" + str(path)


def safe_path(rel: str | None) -> Path:
    """Map request path to filesystem path.

    ALLOW_OUTSIDE_ROOT=no (默认):
      仅允许 DEFAULT_ROOT 及其子目录；越界/.. 一律拒绝。
    ALLOW_OUTSIDE_ROOT=yes:
      允许访问根目录以外的绝对路径（进程有读权限即可）。
    """
    root = ANDROID_LAN_ROOT
    if not root.is_dir():
        abort(503, description=f"browse root not found: {root}")

    raw = (rel or "").replace("\\", "/").strip()
    if raw in ("", "."):
        return root

    if "\x00" in raw:
        abort(400, description="invalid path")

    # Absolute path form: @/abs/path or /abs/path
    if raw.startswith("@/"):
        abs_part = raw[1:]
        if not ALLOW_OUTSIDE_ROOT:
            abort(
                403,
                description=(
                    "outside access disabled. "
                    "Set ALLOW_OUTSIDE_ROOT=yes in android-lan-file-server.conf "
                    "or start with --allow-outside"
                ),
            )
        return _resolve_abs_path(abs_part)

    if raw.startswith("/") and not raw.startswith("//"):
        cand_try = Path(raw).expanduser().resolve()
        root_str = str(root)
        inside = (
            str(cand_try) == root_str
            or str(cand_try).startswith(root_str + os.sep)
        )
        if inside:
            return cand_try
        if not ALLOW_OUTSIDE_ROOT:
            abort(
                403,
                description=(
                    "outside access disabled. "
                    "Set ALLOW_OUTSIDE_ROOT=yes in android-lan-file-server.conf "
                    "or start with --allow-outside"
                ),
            )
        return _resolve_abs_path(raw)

    parts = [p for p in raw.strip("/").split("/") if p not in ("", ".")]
    if any(p == ".." for p in parts):
        abort(400, description="path traversal forbidden")

    candidate = (root / "/".join(parts)).resolve()
    root_str = str(root)
    if str(candidate) != root_str and not str(candidate).startswith(root_str + os.sep):
        if not ALLOW_OUTSIDE_ROOT:
            abort(
                403,
                description=(
                    "outside configured root directory. "
                    "Only files under the default directory and its "
                    "subdirectories are accessible."
                ),
            )
    return candidate


def _resolve_abs_path(abs_path: str) -> Path:
    if "\x00" in abs_path:
        abort(400, description="invalid path")
    candidate = Path(abs_path).expanduser().resolve()
    if not candidate.exists():
        abort(404, description=f"path not found: {candidate}")
    return candidate


def mime_for(path: Path) -> str:
    ext = _ext(path)
    if ext in MIME_OVERRIDES:
        return MIME_OVERRIDES[ext]
    guess, _ = mimetypes.guess_type(str(path))
    return guess or "application/octet-stream"


# ---------------------------------------------------------------------------
# PPT / Office conversion
# ---------------------------------------------------------------------------

# 转换参数变更时递增，作废旧参数生成的缓存
_PDF_CACHE_VER = 4


def _pdf_cache_key(path: Path) -> Path:
    st = path.stat()
    key = f"v{_PDF_CACHE_VER}-{path.name}-{int(st.st_mtime)}-{st.st_size}"
    safe = re.sub(r"[^\w.-]+", "_", key)[:120]
    return CACHE_DIR / safe


def _xlsx_for_preview(src: Path, tmp: Path) -> Path:
    """复制 xlsx 并显式清空页眉页脚。

    文档未声明 headerFooter 时，LibreOffice 会套用默认页面样式，
    把工作表名打进页眉、页码打进页脚（预览里表现为漂浮的孤立字符）。
    此处显式写入空 <headerFooter/> 覆盖该默认行为；只影响预览用的副本，
    原文件不动。
    """
    out = tmp / src.name
    try:
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(
            out, "w", zipfile.ZIP_DEFLATED
        ) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if (
                    item.filename.startswith("xl/worksheets/")
                    and item.filename.endswith(".xml")
                ):
                    text = data.decode("utf-8", errors="replace")
                    text = re.sub(
                        r"<headerFooter\b[^>]*>.*?</headerFooter>", "", text, flags=re.S
                    )
                    text = re.sub(r"<headerFooter\b[^>]*/>", "", text)
                    m = re.search(r"<pageSetup\b[^>]*/>", text)
                    if m:
                        text = text[: m.end()] + "<headerFooter/>" + text[m.end() :]
                    elif "<rowBreaks" in text:
                        text = text.replace(
                            "<rowBreaks", "<headerFooter/><rowBreaks", 1
                        )
                    else:
                        text = text.replace(
                            "</worksheet>", "<headerFooter/></worksheet>", 1
                        )
                    data = text.encode("utf-8")
                zout.writestr(item, data)
        return out
    except (OSError, zipfile.BadZipFile):
        return src


def convert_to_pdf(src: Path, timeout: int = 120) -> Path | None:
    if not SOFFICE:
        return None
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    cache = _pdf_cache_key(src)
    out = cache / f"{src.stem}.pdf"
    if out.is_file() and out.stat().st_size > 0:
        return out
    if out.is_file():
        out.unlink(missing_ok=True)

    with tempfile.TemporaryDirectory(prefix="office-conv-") as tmp_dir:
        tmp_path = Path(tmp_dir)
        conv_src = _xlsx_for_preview(src, tmp_path) if _ext(src) == ".xlsx" else src
        cmd = [
            SOFFICE,
            "--headless",
            "--norestore",
            "--nolockcheck",
            "--nodefault",
            "--nologo",
            "--convert-to",
            "pdf",
            "--outdir",
            tmp_path,
            str(conv_src),
        ]
        try:
            proc = subprocess.run(
                cmd,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=timeout,
                check=False,
            )
        except (subprocess.TimeoutExpired, OSError):
            return None
        produced = Path(tmp_path) / f"{src.stem}.pdf"
        if proc.returncode != 0 or not produced.is_file() or produced.stat().st_size == 0:
            return None
        cache.mkdir(parents=True, exist_ok=True)
        tmp_out = cache / f"{out.name}.tmp.{os.getpid()}"
        shutil.copy2(produced, tmp_out)
        os.replace(tmp_out, out)
    return out


def extract_ppt_text(path: Path, max_slides: int = 30) -> list[dict]:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        return []

    slides: list[dict] = []
    try:
        prs = Presentation(str(path))
    except Exception:
        return []

    for i, slide in enumerate(prs.slides, start=1):
        if i > max_slides:
            break
        texts: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for para in shape.text_frame.paragraphs:
                t = "".join(run.text for run in para.runs).strip()
                if t:
                    texts.append(t)
        title = texts[0][:80] if texts else ""
        bullets = texts[1:40] if len(texts) > 1 else texts
        slides.append({"index": i, "title": title, "bullets": bullets[:30]})
    return slides


def ppt_source_for_display(path: Path) -> tuple[str, Path | None, list[dict]]:
    pdf = convert_to_pdf(path)
    if pdf and pdf.is_file() and pdf.stat().st_size > 0:
        return "pdf", pdf, []
    slides = extract_ppt_text(path)
    if slides:
        return "text", None, slides
    return "none", None, []


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------

@app.route("/")
def index():
    return render_template(
        "index.html",
        root_name=ANDROID_LAN_ROOT.name or "sdcard",
        root_path=str(ANDROID_LAN_ROOT),
        soffice=bool(SOFFICE),
        allow_outside=ALLOW_OUTSIDE_ROOT,
        version=APP_VERSION,
        start_path=request.args.get("path", ""),
    )


@app.route("/api/ls/")
@app.route("/api/ls")
def api_ls():
    rel = request.args.get("path", "")
    target = safe_path(rel)
    if not target.exists():
        return jsonify({"ok": False, "error": "not found", "path": rel}), 404
    if not target.is_dir():
        return jsonify({"ok": False, "error": "not a directory", "path": rel}), 400

    try:
        entries = list(os.scandir(target))
    except PermissionError:
        return jsonify({"ok": False, "error": "permission denied", "path": rel}), 403
    except OSError as e:
        return jsonify({"ok": False, "error": str(e), "path": rel}), 500

    dirs: list[dict] = []
    files: list[dict] = []
    hidden = request.args.get("hidden") == "1"

    for entry in entries:
        name = entry.name
        if not hidden and name.startswith("."):
            continue
        try:
            st = entry.stat(follow_symlinks=False)
            is_dir = entry.is_dir(follow_symlinks=False)
        except OSError:
            continue

        # child path: if parent request was absolute/@, keep that form
        if rel.startswith("@"):
            child_rel = f"{rel.rstrip('/')}/{name}"
        elif rel.startswith("/") and not rel.startswith("//"):
            child_rel = f"{rel.rstrip('/')}/{name}"
        elif rel:
            child_rel = f"{rel.rstrip('/')}/{name}"
        else:
            child_rel = name

        item = {
            "name": name,
            "path": child_rel,
            "is_dir": is_dir,
            "size": st.st_size if not is_dir else 0,
            "mtime": st.st_mtime,
            "size_h": _fmt_size(st.st_size) if not is_dir else "",
            "mtime_h": _fmt_time(st.st_mtime),
        }
        if not is_dir:
            kind = detect_kind(Path(entry.path))
            item["kind"] = kind
            item["previewable"] = kind in PREVIEW_KINDS
            item["ext"] = _ext(Path(name))
        if is_dir:
            dirs.append(item)
        else:
            files.append(item)

    dirs.sort(key=lambda x: x["name"].lower())
    files.sort(key=lambda x: x["name"].lower())

    truncated = False
    if len(dirs) + len(files) > MAX_LIST_ENTRIES:
        truncated = True
        dirs = dirs[: MAX_LIST_ENTRIES // 2]
        files = files[: MAX_LIST_ENTRIES - len(dirs)]

    crumbs = [{"name": ANDROID_LAN_ROOT.name or "sdcard", "path": ""}]
    if rel:
        if rel.startswith("@/") or (rel.startswith("/") and not rel.startswith("//")):
            # absolute browse: show absolute crumb
            abs_p = rel[1:] if rel.startswith("@") else rel
            crumbs = [{"name": abs_p or "/", "path": rel}]
            # still add children crumbs by walking parts under current
        else:
            acc = ""
            for part in rel.split("/"):
                if not part:
                    continue
                acc = f"{acc}/{part}".strip("/") if acc else part
                crumbs.append({"name": part, "path": acc})

    parent = None
    if rel:
        if rel.startswith("@/"):
            parent = "@" + str(Path(rel[1:]).parent)
            if parent == "@":
                parent = ""
        elif rel.startswith("/") and not rel.startswith("//"):
            parent = str(Path(rel).parent)
            if parent == "/":
                parent = ""
        else:
            parts = rel.split("/")
            parent = "/".join(parts[:-1]) if len(parts) > 1 else ""

    return jsonify({
        "ok": True,
        "path": rel,
        "path_abs": str(target),
        "dirs": dirs,
        "files": files,
        "truncated": truncated,
        "count": len(dirs) + len(files),
        "crumbs": crumbs,
        "parent": parent,
        "soffice": bool(SOFFICE),
        "root": str(ANDROID_LAN_ROOT),
        "allow_outside": ALLOW_OUTSIDE_ROOT,
    })


@app.route("/preview/<path:rel>")
def preview(rel: str):
    path = safe_path(rel)
    if not path.exists():
        abort(404, description="file not found")
    if path.is_dir():
        return render_template(
            "index.html",
            root_name=ANDROID_LAN_ROOT.name or "sdcard",
            root_path=str(ANDROID_LAN_ROOT),
            soffice=bool(SOFFICE),
            allow_outside=ALLOW_OUTSIDE_ROOT,
            version=APP_VERSION,
            start_path=rel,
        )

    kind = request.args.get("kind") or detect_kind(path)
    download = request.args.get("download") == "1"
    if download:
        return serve_raw(rel, as_download=True)

    parent_rel = "/".join(rel.split("/")[:-1])
    raw_rel = _rel_for_url(path)
    meta = {
        "name": path.name,
        "path": rel,
        "size": path.stat().st_size,
        "size_h": _fmt_size(path.stat().st_size),
        "mtime_h": _fmt_time(path.stat().st_mtime),
        "kind": kind,
        "ext": _ext(path),
        "mime": mime_for(path),
        "parent": parent_rel,
        "raw_url": f"/raw/{quote(raw_rel, safe='/@')}",
        "download_url": f"/raw/{quote(raw_rel, safe='/@')}?download=1",
        "full_path": str(path),
        "soffice": bool(SOFFICE),
        "allow_outside": ALLOW_OUTSIDE_ROOT,
    }

    if kind == "ppt":
        mode, pdf_path, slides = ppt_source_for_display(path)
        meta["ppt_mode"] = mode
        if mode == "pdf" and pdf_path:
            meta["pdf_url"] = f"/doc-pdf/{quote(raw_rel, safe='/@')}"
            meta["pdf_size_h"] = _fmt_size(pdf_path.stat().st_size)
        meta["slides"] = slides
        return render_template("preview.html", meta=meta)

    if kind == "office":
        pdf = convert_to_pdf(path)
        if pdf and pdf.is_file() and pdf.stat().st_size > 0:
            meta["office_mode"] = "pdf"
            meta["pdf_url"] = f"/doc-pdf/{quote(raw_rel, safe='/@')}"
            meta["pdf_size_h"] = _fmt_size(pdf.stat().st_size)
        else:
            meta["office_mode"] = "none"
        return render_template("preview.html", meta=meta)

    if kind in {"markdown", "text"}:
        try:
            if path.stat().st_size <= MAX_TEXT_BYTES:
                text = path.read_text(encoding="utf-8", errors="replace")
            else:
                with path.open("rb") as f:
                    text = f.read(MAX_TEXT_BYTES).decode("utf-8", errors="replace")
                text += "\n\n... (truncated preview)"
        except OSError as e:
            abort(500, description=str(e))
        meta["content"] = text
        meta["content_url"] = f"/raw/{quote(raw_rel, safe='/@')}?download=0"
        return render_template("preview.html", meta=meta)

    return render_template("preview.html", meta=meta)


@app.route("/doc-pdf/<path:rel>")
def doc_pdf(rel: str):
    path = safe_path(rel)
    if not path.exists() or path.is_dir():
        abort(404)
    if _ext(path) not in (PPT_EXTS | OFFICE_EXTS):
        abort(400, description="not a convertible document")
    pdf = convert_to_pdf(path)
    if not pdf or not pdf.is_file():
        abort(503, description="LibreOffice conversion unavailable or failed")
    return send_file(
        pdf,
        mimetype="application/pdf",
        as_attachment=False,
        download_name=path.stem + ".pdf",
        conditional=True,
        max_age=0,
    )


@app.route("/raw/<path:rel>")
def serve_raw(rel: str, as_download: bool | None = None):
    path = safe_path(rel)
    if not path.exists():
        abort(404, description="file not found")
    if path.is_dir():
        abort(400, description="is a directory")
    if not path.is_file():
        abort(400, description="not a regular file")

    if as_download is None:
        as_download = request.args.get("download") == "1"

    try:
        st = path.stat()
    except OSError as e:
        abort(500, description=str(e))

    resp = send_file(
        path,
        mimetype=mime_for(path),
        as_attachment=as_download,
        download_name=path.name if as_download else None,
        conditional=True,
        max_age=0,
        last_modified=st.st_mtime,
        etag=True,
    )
    resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate"
    resp.headers["Accept-Ranges"] = "bytes"
    return resp


@app.route("/api/kind/<path:rel>")
def api_kind(rel: str):
    path = safe_path(rel)
    if not path.exists():
        return jsonify({"ok": False, "error": "not found"}), 404
    return jsonify({
        "ok": True,
        "kind": detect_kind(path) if path.is_file() else "dir",
        "name": path.name,
        "previewable": path.is_file() and detect_kind(path) in PREVIEW_KINDS,
    })


@app.route("/healthz")
def healthz():
    ok = ANDROID_LAN_ROOT.is_dir() and os.access(ANDROID_LAN_ROOT, os.R_OK)
    return jsonify({
        "ok": ok,
        "root": str(ANDROID_LAN_ROOT),
        "root_readable": ok,
        "allow_outside_root": ALLOW_OUTSIDE_ROOT,
        "conf": str(CONF_PATH),
        "soffice": bool(SOFFICE),
        "version": APP_VERSION,
    }), 200 if ok else 503


@app.errorhandler(404)
def err_404(e):
    if request.path.startswith("/api/"):
        return jsonify({"ok": False, "error": getattr(e, "description", "not found")}), 404
    return render_template("error.html", code=404, message=getattr(e, "description", "Not Found")), 404


@app.errorhandler(400)
def err_400(e):
    if request.path.startswith("/api/"):
        return jsonify({"ok": False, "error": getattr(e, "description", "bad request")}), 400
    return render_template("error.html", code=400, message=getattr(e, "description", "Bad Request")), 400


@app.errorhandler(403)
def err_403(e):
    if request.path.startswith("/api/"):
        return jsonify({"ok": False, "error": getattr(e, "description", "forbidden")}), 403
    return render_template("error.html", code=403, message=getattr(e, "description", "Forbidden")), 403


@app.errorhandler(500)
def err_500(e):
    if request.path.startswith("/api/"):
        return jsonify({"ok": False, "error": "internal error"}), 500
    return render_template("error.html", code=500, message="Internal Server Error"), 500


def main() -> None:
    global ANDROID_LAN_ROOT, ALLOW_OUTSIDE_ROOT
    parser = argparse.ArgumentParser(description="android-lan-file-server local file browser")
    parser.add_argument("--host", default=DEFAULT_HOST)
    parser.add_argument("--port", type=int, default=DEFAULT_PORT)
    parser.add_argument(
        "--root",
        default=str(ANDROID_LAN_ROOT),
        help="browse root directory (overrides android-lan-file-server.conf DEFAULT_ROOT)",
    )
    parser.add_argument(
        "--allow-outside",
        action="store_true",
        help="allow access outside root (overrides ALLOW_OUTSIDE_ROOT=no)",
    )
    parser.add_argument(
        "--deny-outside",
        action="store_true",
        help="force deny access outside root (overrides ALLOW_OUTSIDE_ROOT=yes)",
    )
    parser.add_argument("--debug", action="store_true")
    args = parser.parse_args()

    if args.allow_outside:
        ALLOW_OUTSIDE_ROOT = True
    if args.deny_outside:
        ALLOW_OUTSIDE_ROOT = False

    ANDROID_LAN_ROOT = Path(args.root).expanduser().resolve()
    if not ANDROID_LAN_ROOT.exists():
        raise SystemExit(
            f"[android-lan-file-server] 错误: 根目录不存在: {ANDROID_LAN_ROOT}\n"
            f"请指定一个你确认存在且可访问的目录，例如:\n"
            f"  python3 server.py --root /sdcard/Download"
        )
    if not ANDROID_LAN_ROOT.is_dir():
        raise SystemExit(
            f"[android-lan-file-server] 错误: 根路径不是目录: {ANDROID_LAN_ROOT}"
        )
    if not os.access(ANDROID_LAN_ROOT, os.R_OK | os.X_OK):
        raise SystemExit(
            f"[android-lan-file-server] 错误: 根目录不可读: {ANDROID_LAN_ROOT}\n"
            f"可能原因: 目录不存在于 proot、安卓未授权、或路径拼写错误。"
        )
    try:
        with concurrent.futures.ThreadPoolExecutor(max_workers=1) as ex:
            fut = ex.submit(os.listdir, str(ANDROID_LAN_ROOT))
            fut.result(timeout=3)
    except Exception as e:
        raise SystemExit(
            f"[android-lan-file-server] 错误: 无法列出目录 {ANDROID_LAN_ROOT}: {e}"
        )

    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    port = args.port
    host = args.host
    print(f"[android-lan-file-server] root={ANDROID_LAN_ROOT}")
    print(f"[android-lan-file-server] allow_outside_root={'yes' if ALLOW_OUTSIDE_ROOT else 'no'}")
    print(f"[android-lan-file-server] conf={CONF_PATH}")
    print(f"[android-lan-file-server] listen=http://{host}:{port} (bind={host})")
    print(f"[android-lan-file-server] access_url_loopback=http://127.0.0.1:{port}/")
    # 局域网地址
    lan = []
    try:
        import socket
        # getaddrinfo often returns phone wlan IP
        infos = socket.getaddrinfo(socket.gethostname(), None, socket.AF_INET)
        for info in infos:
            ip = info[4][0]
            if ip and not ip.startswith("127."):
                if ip not in lan:
                    lan.append(ip)
    except OSError:
        pass
    if not lan:
        try:
            import subprocess as _sp
            out = _sp.check_output(
                ["hostname", "-I"], text=True, stderr=_sp.DEVNULL
            )
            for tok in out.split():
                if tok and not tok.startswith("127.") and "." in tok:
                    lan.append(tok)
        except Exception:
            pass
    if lan:
        print("[android-lan-file-server] access_url_lan:")
        for ip in lan:
            print(f"[android-lan-file-server]   http://{ip}:{port}/")
        if host == "0.0.0.0":
            print("[android-lan-file-server]   (监听 0.0.0.0，局域网他人可直接打开上述地址)")
        else:
            print(
                "[android-lan-file-server]   当前绑定非 0.0.0.0，局域网可能无法连接；"
                "如需分享请将 ANDROID_LAN_HOST 设为 0.0.0.0"
            )
    else:
        print(f"[android-lan-file-server] access_url_lan=http://<本机IP>:{port}/")
    # 自写 PID，供 stop.sh 与启动横幅展示（前台/后台两种模式都生效）
    try:
        Path("/tmp/android-lan-file-server.pid").write_text(str(os.getpid()), encoding="utf-8")
    except OSError:
        pass
    app.run(host=host, port=port, debug=args.debug, threaded=True)


if __name__ == "__main__":
    main()
